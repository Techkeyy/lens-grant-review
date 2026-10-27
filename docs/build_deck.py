"""Build the editable Lens hackathon slide deck (Python package: python-pptx)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Lens_Open_Agent_Hackathon_2026.pptx"
NAVY = RGBColor(18, 35, 49)
INK = RGBColor(31, 52, 63)
TEAL = RGBColor(0, 119, 113)
PALE = RGBColor(236, 246, 243)
WHITE = RGBColor(255, 255, 255)
CREAM = RGBColor(249, 248, 242)
GRAY = RGBColor(89, 107, 114)
AMBER = RGBColor(180, 96, 42)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, color, radius=False):
    s = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    if radius:
        s.adjustments[0] = 0.12
    return s


def txt(slide, value, x, y, w, h, size=20, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(7)
        if align:
            p.alignment = align
    return box


def base(title, number, subtitle=None):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CREAM
    rect(slide, 0, 0, 0.12, 7.5, TEAL)
    txt(slide, title, 0.58, 0.43, 12.1, 0.7, 29, NAVY, True)
    if subtitle:
        txt(slide, subtitle, 0.6, 1.15, 12.0, 0.6, 14, GRAY)
    txt(slide, f"Lens   /   Open Agent Hackathon 2026   /   {number:02}", 0.6, 7.12, 12, 0.23, 10, GRAY)
    return slide


def panel(slide, x, y, w, h, title, body, accent=TEAL):
    rect(slide, x, y, w, h, WHITE, True)
    rect(slide, x, y, 0.08, h, accent)
    txt(slide, title, x + 0.28, y + 0.25, w - 0.55, 0.62, 21, NAVY, True)
    txt(slide, body, x + 0.28, y + 0.95, w - 0.55, h - 1.15, 17, INK)


# 1: concrete problem; values are from the synthetic sample only.
s = base("Two reports can claim the same four hours", 1,
         "A small charity must check funder reports before they leave. Synthetic demonstration data.")
panel(s, 0.7, 2.0, 3.55, 3.1, "Reading Futures", "Timesheet TS-101\n4 hours claimed")
panel(s, 5.0, 2.0, 3.55, 3.1, "Community Learning", "The same TS-101\n4 hours claimed")
panel(s, 9.25, 2.0, 3.35, 3.1, "Original timesheet", "TS-101 records\n4 hours total", AMBER)
txt(s, "8 hours claimed from one 4-hour record", 1.0, 5.58, 11.3, 0.6, 24, NAVY, True, PP_ALIGN.CENTER)

# 2: product outcome.
s = base("Lens shows the records behind the question", 2,
         "A review workspace for the person preparing two funder reports.")
panel(s, 0.7, 2.0, 3.85, 3.7, "1  Bring records together", "Two draft report CSVs\nTimesheets, receipts and activity log")
panel(s, 4.75, 2.0, 3.85, 3.7, "2  Check the claims", "See both report rows, the source row and the calculation")
panel(s, 8.8, 2.0, 3.85, 3.7, "3  Decide and rerun", "Correct an allocation or write an explanation; review the result again")
txt(s, "Lens raises a question for a reviewer. It does not certify compliance or submit reports.",
    0.76, 6.22, 11.9, 0.45, 15, GRAY)

# 3: actual product capture from the current release, with the synthetic label.
s = base("The result puts the question before the raw rows", 3,
         "Actual Lens results screen, synthetic data. A reviewer can inspect evidence and act.")
s.shapes.add_picture(str(ROOT / "images" / "results-desktop.png"), Inches(0.72), Inches(1.86), height=Inches(4.45))
panel(s, 8.25, 1.86, 4.38, 4.45, "Review path",
      "Sign in and open a case.\n\nCheck both reports.\n\nRead the 8-versus-4 question.\n\nCorrect, rerun and export.")

# 4: central evidence card.
s = base("The answer includes its source and arithmetic", 4,
         "Lens groups allocations by stable source ID across both reports.")
panel(s, 0.68, 1.94, 3.8, 3.55, "Report A", "TS-101\n4 staff hours\nReading Futures")
panel(s, 4.75, 1.94, 3.8, 3.55, "Report B", "TS-101\n4 staff hours\nCommunity Learning")
panel(s, 8.82, 1.94, 3.8, 3.55, "Evidence", "TS-101\n4 hours on the original timesheet", AMBER)
rect(s, 2.35, 5.85, 8.65, 0.76, PALE, True)
txt(s, "4 + 4 = 8 hours claimed  >  4 hours recorded", 2.55, 5.98, 8.25, 0.48, 22, NAVY, True, PP_ALIGN.CENTER)

# 5: current architecture, separating verified evidence from provider readiness.
s = base("Evidence and model decisions have separate jobs", 5,
         "Browser UI + Python HTTP server + Supabase saved reviews; no model may rewrite a finding.")
for x, label, body in [
    (0.65, "Zetaris data", "Read five registered grant tables; retain source keys, hashes and query IDs."),
    (4.43, "Lens + Nightshift", "Lens checks arithmetic. A configured Nightshift model prioritizes one bounded follow-up; live proof pending."),
    (8.21, "World Model", "Meterless reference runtime links source, finding and reviewer events. Supabase saves public reviews."),
]:
    panel(s, x, 2.15, 3.55, 3.4, label, body)
txt(s, "Human correction changes the saved review; the original source remains available.",
    1.0, 6.0, 11.2, 0.54, 19, NAVY, True, PP_ALIGN.CENTER)

# 6: sponsor contribution and honest boundaries.
s = base("Sponsor contribution: proof and remaining gate", 6,
         "The product combines identified source queries with an event history and a bounded model route.")
panel(s, 0.6, 1.93, 3.92, 4.55, "Zetaris", "Five registered synthetic grant tables were queried live. Source keys, row counts, snapshot hashes and request IDs stay with the review.")
panel(s, 4.72, 1.93, 3.92, 4.55, "Meterless Engine", "World Model reference runtime stores linked sources, findings and reviewer actions. A fresh process can rebuild the history.")
panel(s, 8.84, 1.93, 3.92, 4.55, "Nightshift", "The provider adapter asks a model to prioritize an open question. A real hotspot request and response remain unverified.", AMBER)

# 7: scope, roadmap and close.
s = base("Give the reviewer a chance to fix the report", 7,
         "Built for a focused review of two grants and one reporting period.")
panel(s, 0.73, 2.0, 5.84, 3.88, "Current scope", "Five structured CSV inputs or owner-controlled Zetaris tables; cross-grant allocation checks; activity gaps; Google sign-in; saved reviews; JSON export.")
panel(s, 6.77, 2.0, 5.84, 3.88, "Next steps", "Funder templates and document imports; accounting connectors; grant-specific reporting terms; collaboration and broader real-record access controls.", AMBER)
txt(s, "Before a report goes out, see which claims the records support and which need attention.",
    0.86, 6.17, 11.8, 0.6, 20, NAVY, True, PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
